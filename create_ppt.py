from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
import copy

prs = Presentation()
prs.slide_width = Inches(13.33)
prs.slide_height = Inches(7.5)

# ── Colour palette ──────────────────────────────────────────────
YELLOW    = RGBColor(0xFF, 0xD7, 0x00)   # Bollywood gold
ORANGE    = RGBColor(0xFF, 0x6B, 0x00)
DARK      = RGBColor(0x1A, 0x1A, 0x2E)   # deep navy
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_BG  = RGBColor(0xF5, 0xF0, 0xE8)   # warm off-white
GREEN     = RGBColor(0x06, 0xD6, 0x72)
RED       = RGBColor(0xEF, 0x23, 0x3C)
PURPLE    = RGBColor(0x8B, 0x5C, 0xF6)
BLUE      = RGBColor(0x00, 0x8B, 0xFF)

blank_layout = prs.slide_layouts[6]  # completely blank

# ── Helper utilities ────────────────────────────────────────────

def bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color

def box(slide, l, t, w, h, bg_color=None, border_color=None):
    shape = slide.shapes.add_shape(1, Inches(l), Inches(t), Inches(w), Inches(h))
    shape.line.fill.background() if border_color is None else None
    if bg_color:
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
    else:
        shape.fill.background()
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(2)
    else:
        shape.line.fill.background()
    return shape

def txt(slide, text, l, t, w, h,
        size=24, bold=False, color=WHITE, align=PP_ALIGN.LEFT,
        italic=False, wrap=True):
    txb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    txb.word_wrap = wrap
    tf = txb.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return txb

def multiline(slide, lines, l, t, w, h,
              size=20, bold=False, color=WHITE, align=PP_ALIGN.LEFT,
              line_spacing=1.2, italic=False):
    txb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    txb.word_wrap = True
    tf = txb.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(4)
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = color

def pill(slide, label, l, t, w=2.2, h=0.45, bg_color=ORANGE, text_color=WHITE, size=16):
    b = box(slide, l, t, w, h, bg_color=bg_color)
    txt(slide, label, l+0.1, t+0.03, w-0.1, h-0.05,
        size=size, bold=True, color=text_color, align=PP_ALIGN.CENTER)

def divider(slide, y, color=YELLOW, width=12, left=0.65):
    ln = slide.shapes.add_shape(1, Inches(left), Inches(y), Inches(width), Inches(0.04))
    ln.fill.solid()
    ln.fill.fore_color.rgb = color
    ln.line.fill.background()

# ════════════════════════════════════════════════════════════════
# SLIDE 1 — HOOK (Baburao opener)
# ════════════════════════════════════════════════════════════════
s1 = prs.slides.add_slide(blank_layout)
bg(s1, DARK)

# big yellow quote box
box(s1, 0.5, 0.6, 8.2, 2.8, bg_color=YELLOW)
txt(s1, '"Yeh AI kya hota hai??"', 0.7, 0.75, 7.8, 1.2,
    size=48, bold=True, color=DARK, align=PP_ALIGN.CENTER)
txt(s1, '— Baburao, every colleague, 2024', 0.7, 1.85, 7.8, 0.8,
    size=22, bold=False, color=DARK, align=PP_ALIGN.CENTER, italic=True)

# right side label
box(s1, 9.2, 0.6, 3.6, 2.8, bg_color=ORANGE)
txt(s1, '🎬', 9.4, 0.8, 3.2, 1.0, size=60, align=PP_ALIGN.CENTER, color=WHITE)
txt(s1, 'Hera Pheri\nVibes', 9.4, 1.7, 3.2, 1.2,
    size=24, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

divider(s1, 3.7)

txt(s1, 'But here\'s the thing —', 0.7, 3.9, 12, 0.7,
    size=28, color=WHITE, align=PP_ALIGN.LEFT)
txt(s1, 'AI is not the problem.   Your prompt is.', 0.7, 4.5, 12, 0.9,
    size=38, bold=True, color=YELLOW, align=PP_ALIGN.LEFT)

txt(s1, 'Today we fix that. In 20 minutes. 🚀', 0.7, 5.6, 12, 0.7,
    size=22, color=WHITE, align=PP_ALIGN.LEFT, italic=True)

pill(s1, '⏱ 20 MIN SESSION', 9.5, 6.6, 3.3, 0.5, bg_color=GREEN, size=15)

# ════════════════════════════════════════════════════════════════
# SLIDE 2 — INTERN ANALOGY (split)
# ════════════════════════════════════════════════════════════════
s2 = prs.slides.add_slide(blank_layout)
bg(s2, DARK)

txt(s2, 'AI = Your New Intern', 0.5, 0.2, 12.3, 0.8,
    size=38, bold=True, color=YELLOW, align=PP_ALIGN.CENTER)
divider(s2, 1.1)

# LEFT panel — Day 1 intern
box(s2, 0.4, 1.3, 5.8, 5.5, bg_color=RED)
txt(s2, '😰  Intern — Day 1', 0.6, 1.45, 5.4, 0.7,
    size=24, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
divider(s2, 2.15, color=WHITE, width=5.4, left=0.6)
multiline(s2,
    ['You said:  "Write me an email"',
     '',
     'Intern writes a 3-page essay.',
     'Wrong tone. Wrong person.',
     'Mentions things you never said.',
     '',
     '😤 You redo it yourself.'],
    0.6, 2.3, 5.4, 4.0,
    size=19, color=WHITE)

# RIGHT panel — After good brief
box(s2, 6.8, 1.3, 5.8, 5.5, bg_color=GREEN)
txt(s2, '🚀  Intern — After a Good Brief', 6.95, 1.45, 5.5, 0.7,
    size=24, bold=True, color=DARK, align=PP_ALIGN.CENTER)
divider(s2, 2.15, color=DARK, width=5.4, left=6.95)
multiline(s2,
    ['You said:  "You\'re a biz writer.',
     'Write a 80-word follow-up to',
     'Rajesh from TechCorp. Meeting',
     'was about Q3 budget. Formal tone."',
     '',
     'Intern nails it. First try.',
     '',
     '🎯 You just saved 20 minutes.'],
    6.95, 2.3, 5.4, 4.0,
    size=19, color=DARK)

# centre arrow
txt(s2, '➡', 6.0, 3.7, 0.8, 0.8, size=40, color=YELLOW, align=PP_ALIGN.CENTER)

# ════════════════════════════════════════════════════════════════
# SLIDE 3 — THE PROBLEM (bad prompts)
# ════════════════════════════════════════════════════════════════
s3 = prs.slides.add_slide(blank_layout)
bg(s3, DARK)

txt(s3, 'Why AI Gives You Garbage Sometimes', 0.5, 0.2, 12.3, 0.75,
    size=36, bold=True, color=RED, align=PP_ALIGN.CENTER)
divider(s3, 1.1)

txt(s3, '(Hint: it\'s not the AI 😬)', 0.5, 1.15, 12.3, 0.5,
    size=20, color=WHITE, align=PP_ALIGN.CENTER, italic=True)

# 3 bad prompt cards
bad = [
    ('❌  Too Vague', '"Write me an email"', 'No recipient, no context,\nno tone. AI guesses everything.'),
    ('❌  No Context', '"Summarise this report"', 'Who is reading it? What do\nthey care about? Unknown.'),
    ('❌  No Format', '"Make me a formula"', 'For what data? What output?\nWhat version of Excel?'),
]
for i, (title, prompt, explain) in enumerate(bad):
    x = 0.5 + i * 4.3
    box(s3, x, 1.8, 4.0, 4.8, bg_color=RGBColor(0x3A, 0x0A, 0x0A))
    box(s3, x, 1.8, 4.0, 0.7, bg_color=RED)
    txt(s3, title, x+0.1, 1.85, 3.8, 0.6, size=18, bold=True, color=WHITE)
    box(s3, x+0.1, 2.6, 3.8, 0.85, bg_color=RGBColor(0x2A, 0x2A, 0x2A), border_color=RED)
    txt(s3, prompt, x+0.2, 2.65, 3.6, 0.75, size=17, color=YELLOW, italic=True)
    txt(s3, explain, x+0.1, 3.6, 3.8, 1.8, size=17, color=WHITE)

txt(s3, '🎬 Classic Baburao energy. Confused. Getting nowhere.', 0.5, 6.8, 12.3, 0.5,
    size=18, color=YELLOW, align=PP_ALIGN.CENTER, italic=True)

# ════════════════════════════════════════════════════════════════
# SLIDE 4 — THE FORMULA (5 parts)
# ════════════════════════════════════════════════════════════════
s4 = prs.slides.add_slide(blank_layout)
bg(s4, DARK)

txt(s4, 'The 5-Part Formula  🎓  (Rancho Mode: ON)', 0.5, 0.15, 12.3, 0.75,
    size=34, bold=True, color=YELLOW, align=PP_ALIGN.CENTER)
divider(s4, 1.0)

parts = [
    (BLUE,   '1. ROLE',        'Who should AI be?',      '"You are a senior HR manager..."'),
    (PURPLE, '2. TASK',        'What exactly do you need?', '"Write a 3-bullet summary..."'),
    (ORANGE, '3. CONTEXT',     'What background info?',   '"The audience is non-technical..."'),
    (GREEN,  '4. FORMAT',      'How should output look?', '"Use bullet points, under 100 words"'),
    (RED,    '5. CONSTRAINTS', 'What to avoid?',          '"Don\'t use jargon. Keep it formal."'),
]

for i, (color, label, desc, example) in enumerate(parts):
    y = 1.2 + i * 1.1
    box(s4, 0.4, y, 1.8, 0.85, bg_color=color)
    txt(s4, label, 0.45, y+0.08, 1.7, 0.7, size=16, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    txt(s4, desc, 2.4, y+0.05, 4.2, 0.75, size=18, color=WHITE)
    box(s4, 6.8, y, 6.1, 0.85, bg_color=RGBColor(0x2A, 0x2A, 0x2A), border_color=color)
    txt(s4, example, 6.95, y+0.1, 5.8, 0.65, size=17, color=YELLOW, italic=True)

txt(s4, '👆 This is your new superpower. Screenshot it.', 0.5, 6.8, 12.3, 0.5,
    size=18, color=GREEN, align=PP_ALIGN.CENTER, bold=True)

# ════════════════════════════════════════════════════════════════
# SLIDE 5 — USE CASE: EMAIL
# ════════════════════════════════════════════════════════════════
s5 = prs.slides.add_slide(blank_layout)
bg(s5, DARK)

pill(s5, '📧  USE CASE 1: EMAIL', 0.4, 0.15, 4.5, 0.6, bg_color=BLUE)
txt(s5, 'Before vs After', 5.2, 0.2, 7.5, 0.55,
    size=22, color=WHITE, italic=True)
divider(s5, 0.95)

box(s5, 0.4, 1.1, 5.8, 5.6, bg_color=RGBColor(0x1E, 0x10, 0x10))
box(s5, 0.4, 1.1, 5.8, 0.6, bg_color=RED)
txt(s5, '❌  Bad Prompt', 0.6, 1.15, 5.4, 0.5, size=20, bold=True, color=WHITE)
txt(s5, '"Write me a follow up email"', 0.6, 1.85, 5.4, 0.7,
    size=18, color=YELLOW, italic=True)
divider(s5, 2.65, color=RED, width=5.4, left=0.6)
txt(s5, 'AI Output:', 0.6, 2.75, 5.4, 0.4, size=15, color=RGBColor(0xAA,0xAA,0xAA))
multiline(s5,
    ['"Dear Sir/Madam,',
     'I hope this email finds you well.',
     'I am writing to follow up on our',
     'recent conversation...',
     '',
     '❌ Generic. Could be anyone.',
     '❌ Wrong tone. Wrong length.'],
    0.6, 3.1, 5.4, 3.4, size=16, color=WHITE)

box(s5, 6.9, 1.1, 5.9, 5.6, bg_color=RGBColor(0x05, 0x1A, 0x10))
box(s5, 6.9, 1.1, 5.9, 0.6, bg_color=GREEN)
txt(s5, '✅  Good Prompt', 7.1, 1.15, 5.5, 0.5, size=20, bold=True, color=DARK)
multiline(s5,
    ['"You are a business writer.',
     'Write a 80-word follow-up email',
     'to Rajesh from TechCorp after a',
     'Q3 budget meeting. Formal tone.',
     'End with a next-step ask."'],
    7.05, 1.85, 5.6, 1.8, size=16, color=YELLOW, italic=True)
divider(s5, 3.75, color=GREEN, width=5.6, left=7.05)
txt(s5, 'AI Output:', 7.1, 3.85, 5.5, 0.4, size=15, color=RGBColor(0xAA,0xAA,0xAA))
multiline(s5,
    ['"Hi Rajesh, great connecting today.',
     'As discussed, Q3 budget review is',
     'key for us. Would Thursday work',
     'for a 30-min call to align?..."',
     '',
     '✅ Right person. Right tone.',
     '✅ Correct length. Has a CTA.'],
    7.1, 4.2, 5.6, 2.4, size=16, color=WHITE)

txt(s5, '⏱ Same AI. Different prompt. Totally different result.', 0.4, 6.85, 12.5, 0.45,
    size=17, color=YELLOW, align=PP_ALIGN.CENTER, italic=True)

# ════════════════════════════════════════════════════════════════
# SLIDE 6 — USE CASE: SUMMARISING DOCS
# ════════════════════════════════════════════════════════════════
s6 = prs.slides.add_slide(blank_layout)
bg(s6, DARK)

pill(s6, '📄  USE CASE 2: SUMMARISING DOCS', 0.4, 0.15, 5.8, 0.6, bg_color=PURPLE)
divider(s6, 0.95)

box(s6, 0.4, 1.1, 5.8, 5.6, bg_color=RGBColor(0x15, 0x0E, 0x2A))
box(s6, 0.4, 1.1, 5.8, 0.6, bg_color=RED)
txt(s6, '❌  Bad Prompt', 0.6, 1.15, 5.4, 0.5, size=20, bold=True, color=WHITE)
txt(s6, '"Summarise this"', 0.6, 1.85, 5.4, 0.6,
    size=18, color=YELLOW, italic=True)
divider(s6, 2.6, color=RED, width=5.4, left=0.6)
txt(s6, 'AI Output:', 0.6, 2.7, 5.4, 0.4, size=15, color=RGBColor(0xAA,0xAA,0xAA))
multiline(s6,
    ['"This report discusses various',
     'aspects of the company including',
     'financials, operations, HR and...',
     '',
     '❌ Covers everything.',
     '❌ Useful to no one.',
     '❌ You still have to read it.'],
    0.6, 3.05, 5.4, 3.5, size=16, color=WHITE)

box(s6, 6.9, 1.1, 5.9, 5.6, bg_color=RGBColor(0x08, 0x15, 0x20))
box(s6, 6.9, 1.1, 5.9, 0.6, bg_color=GREEN)
txt(s6, '✅  Good Prompt', 7.1, 1.15, 5.5, 0.5, size=20, bold=True, color=DARK)
multiline(s6,
    ['"Summarise this report for a',
     'senior manager who has 2 mins.',
     'Pull out: top 3 risks, key',
     'decision needed, and any',
     'financial red flags. Bullets."'],
    7.05, 1.85, 5.6, 2.0, size=16, color=YELLOW, italic=True)
divider(s6, 3.95, color=GREEN, width=5.6, left=7.05)
txt(s6, 'AI Output:', 7.1, 4.05, 5.5, 0.4, size=15, color=RGBColor(0xAA,0xAA,0xAA))
multiline(s6,
    ['"• Risk 1: Supply chain delay',
     '• Risk 2: Budget overrun (₹2Cr)',
     '• Risk 3: Attrition in Q4',
     '• Decision needed: Vendor lock-in',
     '',
     '✅ Scannable. Decision-ready.',
     '✅ Manager loves you now. 😎'],
    7.1, 4.4, 5.6, 2.4, size=16, color=WHITE)

txt(s6, '🎬 Rancho energy — simple brief, brilliant output.', 0.4, 6.85, 12.5, 0.45,
    size=17, color=YELLOW, align=PP_ALIGN.CENTER, italic=True)

# ════════════════════════════════════════════════════════════════
# SLIDE 7 — USE CASE: EXCEL
# ════════════════════════════════════════════════════════════════
s7 = prs.slides.add_slide(blank_layout)
bg(s7, DARK)

pill(s7, '📊  USE CASE 3: EXCEL FORMULAS', 0.4, 0.15, 5.5, 0.6, bg_color=GREEN, text_color=DARK)
divider(s7, 0.95)

box(s7, 0.4, 1.1, 5.8, 5.6, bg_color=RGBColor(0x10, 0x1A, 0x10))
box(s7, 0.4, 1.1, 5.8, 0.6, bg_color=RED)
txt(s7, '❌  Bad Prompt', 0.6, 1.15, 5.4, 0.5, size=20, bold=True, color=WHITE)
txt(s7, '"Give me an Excel formula"', 0.6, 1.85, 5.4, 0.6,
    size=18, color=YELLOW, italic=True)
divider(s7, 2.6, color=RED, width=5.4, left=0.6)
txt(s7, 'AI Output:', 0.6, 2.7, 5.4, 0.4, size=15, color=RGBColor(0xAA,0xAA,0xAA))
multiline(s7,
    ['"Sure! Here are some common',
     'Excel formulas: SUM, AVERAGE,',
     'VLOOKUP, IF, COUNTIF..."',
     '',
     '❌ A tutorial. Not an answer.',
     '❌ You still have to figure',
     '    out which one to use.'],
    0.6, 3.05, 5.4, 3.5, size=16, color=WHITE)

box(s7, 6.9, 1.1, 5.9, 5.6, bg_color=RGBColor(0x08, 0x15, 0x10))
box(s7, 6.9, 1.1, 5.9, 0.6, bg_color=GREEN)
txt(s7, '✅  Good Prompt', 7.1, 1.15, 5.5, 0.5, size=20, bold=True, color=DARK)
multiline(s7,
    ['"I have an Excel sheet. Col A =',
     'Employee names, Col B = Sales.',
     'I want to highlight anyone who',
     'sold more than ₹5L in red.',
     'Give me the conditional',
     'formatting formula."'],
    7.05, 1.85, 5.6, 2.2, size=16, color=YELLOW, italic=True)
divider(s7, 4.15, color=GREEN, width=5.6, left=7.05)
txt(s7, 'AI Output:', 7.1, 4.25, 5.5, 0.4, size=15, color=RGBColor(0xAA,0xAA,0xAA))
multiline(s7,
    ['"Select Col B → Conditional',
     'Formatting → New Rule →',
     'Use formula: =B2>500000',
     'Set fill to red. Done."',
     '',
     '✅ Step-by-step. Plug and play.',
     '✅ No guessing. No Googling.'],
    7.1, 4.6, 5.6, 2.2, size=16, color=WHITE)

txt(s7, '💡 Context is everything. AI can\'t read your screen — describe it.', 0.4, 6.85, 12.5, 0.45,
    size=17, color=YELLOW, align=PP_ALIGN.CENTER, italic=True)

# ════════════════════════════════════════════════════════════════
# SLIDE 8 — THE BACK AND FORTH
# ════════════════════════════════════════════════════════════════
s8 = prs.slides.add_slide(blank_layout)
bg(s8, DARK)

txt(s8, 'AI is a Conversation, Not a Vending Machine', 0.5, 0.15, 12.3, 0.75,
    size=34, bold=True, color=YELLOW, align=PP_ALIGN.CENTER)
divider(s8, 1.0)

txt(s8, '"Kabhi Khushi Kabhie Gham — it\'s a journey, not one scene."',
    0.5, 1.1, 12.3, 0.6, size=20, color=WHITE, align=PP_ALIGN.CENTER, italic=True)

steps = [
    (BLUE,   '1. Send your prompt',      'Use the 5-part formula as a starting point'),
    (PURPLE, '2. Review the output',     'Is it close? What\'s off — tone, length, detail?'),
    (ORANGE, '3. Refine in the same chat','Say: "Make it shorter" / "More formal" / "Add X"'),
    (GREEN,  '4. Done in 2-3 rounds',    'Not 10. A good first prompt = fewer rounds needed'),
]

for i, (color, step, desc) in enumerate(steps):
    y = 2.0 + i * 1.1
    box(s8, 0.4, y, 0.7, 0.8, bg_color=color)
    txt(s8, str(i+1), 0.4, y+0.08, 0.7, 0.65, size=30, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    txt(s8, step, 1.3, y+0.05, 4.5, 0.7, size=20, bold=True, color=color)
    txt(s8, desc, 6.0, y+0.1, 6.8, 0.65, size=18, color=WHITE)

box(s8, 0.4, 6.4, 12.5, 0.8, bg_color=RGBColor(0x06, 0x40, 0x30))
txt(s8, '🎯 Pro tip: If output is totally wrong, don\'t re-send the same prompt. Add context.',
    0.6, 6.48, 12.1, 0.6, size=18, color=GREEN, bold=True)

# ════════════════════════════════════════════════════════════════
# SLIDE 9 — COMMON MISTAKES
# ════════════════════════════════════════════════════════════════
s9 = prs.slides.add_slide(blank_layout)
bg(s9, DARK)

txt(s9, 'Quick Hits — Mistakes Everyone Makes 🤦', 0.5, 0.15, 12.3, 0.75,
    size=34, bold=True, color=RED, align=PP_ALIGN.CENTER)
divider(s9, 1.0)

mistakes = [
    ('😬', 'Too vague',              '"Write something" → AI writes anything'),
    ('🤷', 'No audience context',    'AI doesn\'t know if it\'s for a CEO or intern'),
    ('🔁', 'Accepting first output', 'First draft ≠ final answer. Always refine.'),
    ('📋', 'No format specified',    'AI writes an essay when you wanted 3 bullets'),
    ('🙈', 'Pasting raw data blind', 'Describe your data — don\'t just dump it'),
    ('📵', 'Starting fresh each time','Continue the chat. AI has memory within a session'),
]

for i, (emoji, mistake, tip) in enumerate(mistakes):
    row = i // 2
    col = i % 2
    x = 0.4 + col * 6.45
    y = 1.2 + row * 1.7
    box(s9, x, y, 6.1, 1.5, bg_color=RGBColor(0x25, 0x10, 0x10), border_color=RED)
    txt(s9, emoji, x+0.15, y+0.1, 0.8, 1.2, size=32, align=PP_ALIGN.CENTER, color=WHITE)
    txt(s9, mistake, x+1.1, y+0.1, 4.8, 0.6, size=19, bold=True, color=RED)
    txt(s9, tip, x+1.1, y+0.65, 4.8, 0.7, size=15, color=WHITE)

txt(s9, '🎬 Baburao se Rancho bano — these are the differences.', 0.5, 6.85, 12.3, 0.45,
    size=17, color=YELLOW, align=PP_ALIGN.CENTER, italic=True)

# ════════════════════════════════════════════════════════════════
# SLIDE 10 — LIVE EXERCISE
# ════════════════════════════════════════════════════════════════
s10 = prs.slides.add_slide(blank_layout)
bg(s10, DARK)

box(s10, 0, 0, 13.33, 7.5, bg_color=RGBColor(0x0A, 0x20, 0x1A))

txt(s10, '⚡ YOUR TURN', 0.5, 0.2, 12.3, 0.85,
    size=48, bold=True, color=GREEN, align=PP_ALIGN.CENTER)
divider(s10, 1.15, color=GREEN)

txt(s10, 'Rewrite this bad prompt using the 5-part formula:', 0.5, 1.3, 12.3, 0.6,
    size=24, color=WHITE, align=PP_ALIGN.CENTER)

box(s10, 2.0, 2.05, 9.3, 1.0, bg_color=RED)
txt(s10, '"Help me with my presentation"', 2.2, 2.15, 8.9, 0.8,
    size=26, bold=True, color=WHITE, align=PP_ALIGN.CENTER, italic=True)

multiline(s10,
    ['Think about:', '  • Who is AI in this scenario? (Role)',
     '  • What exactly do you need? (Task)',
     '  • What\'s the presentation about? (Context)',
     '  • How long / what format? (Format)',
     '  • Any constraints? (Constraints)'],
    1.5, 3.25, 10.3, 2.8, size=20, color=WHITE)

pill(s10, '⏱ 3 MINUTES — GO!', 4.5, 6.4, 4.3, 0.7,
     bg_color=YELLOW, text_color=DARK, size=20)

# ════════════════════════════════════════════════════════════════
# SLIDE 11 — CHEAT SHEET
# ════════════════════════════════════════════════════════════════
s11 = prs.slides.add_slide(blank_layout)
bg(s11, DARK)

box(s11, 0.3, 0.1, 12.73, 7.3, bg_color=RGBColor(0x12, 0x12, 0x28), border_color=YELLOW)

txt(s11, '📸 SCREENSHOT THIS — The Prompt Cheat Sheet', 0.5, 0.2, 12.3, 0.7,
    size=30, bold=True, color=YELLOW, align=PP_ALIGN.CENTER)
divider(s11, 1.0, color=YELLOW)

parts2 = [
    (BLUE,   '1  ROLE',        '"You are a [job title / expert]..."'),
    (PURPLE, '2  TASK',        '"Write / Summarise / Create / Fix..."'),
    (ORANGE, '3  CONTEXT',     '"Background: [who, what, why]..."'),
    (GREEN,  '4  FORMAT',      '"Respond in [bullets / email / table]..."'),
    (RED,    '5  CONSTRAINTS', '"Keep it under X words / Avoid Y..."'),
]

for i, (color, label, example) in enumerate(parts2):
    y = 1.15 + i * 1.05
    box(s11, 0.5, y, 2.0, 0.85, bg_color=color)
    txt(s11, label, 0.55, y+0.1, 1.9, 0.65, size=17, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    box(s11, 2.65, y, 9.9, 0.85, bg_color=RGBColor(0x1E, 0x1E, 0x35), border_color=color)
    txt(s11, example, 2.8, y+0.12, 9.6, 0.65, size=19, color=YELLOW, italic=True)

box(s11, 0.5, 6.5, 12.33, 0.65, bg_color=GREEN)
txt(s11, '"Aal Izz Well" — Follow the formula and AI will never let you down. 🎓',
    0.6, 6.55, 12.1, 0.55, size=18, bold=True, color=DARK, align=PP_ALIGN.CENTER)

# ════════════════════════════════════════════════════════════════
# Save
# ════════════════════════════════════════════════════════════════
out = "/home/user/jcode/Prompt_Like_A_Pro.pptx"
prs.save(out)
print(f"Saved → {out}")
