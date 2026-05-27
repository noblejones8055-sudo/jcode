#!/usr/bin/env python3
"""Part 1: Sections 5.3 – 5.3.4 of the corrected Pollution Report."""

import sys
sys.path.insert(0, '/home/user/jcode')
from build_report_helpers import *
from docx import Document

def build_part1(doc):

    # ═══════════════════════════════════════════════════════════
    # 5.3  POLLUTION
    # ═══════════════════════════════════════════════════════════
    heading(doc, '5.3  Pollution', level=1)

    # Intro para 1 – ends with inline RULE_02 GRI code
    p1 = doc.add_paragraph()
    txt1 = ("Pollution prevention and environmental stewardship are core commitments within "
            "XYZ Logistics LLC's environmental strategy. The Company recognises that effective "
            "pollution management is fundamental to protecting ecosystems, human health, and the "
            "communities in which it operates. XYZ Logistics LLC maintains a comprehensive "
            "Pollution Prevention and Control programme designed to minimise emissions to air, "
            "water, and soil; responsibly manage substances of concern (SoC) and substances of "
            "very high concern (SVHC); and proactively mitigate pollution-related financial and "
            "operational risks across all levels of the organisation.  ")
    fmt(p1.add_run(txt1), txt1, sz=11)
    disc_r02(p1, 'GRI 305-7')

    # Intro para 2 – ends with inline RULE_02 ESRS codes
    p2 = doc.add_paragraph()
    txt2 = ("This section presents XYZ Logistics LLC's Pollution disclosure for FY 2025, "
            "aligned with the relevant European Sustainability Reporting Standards (ESRS) "
            "and applicable GRI Standards.  ")
    fmt(p2.add_run(txt2), txt2, sz=11)
    disc_r02(p2, 'ESRS E2 | ESRS E2-IRO-1 | ESRS E2-4 | ESRS E2-5 | ESRS E2-6')

    # Overview table (no disclosure code on title – introductory context table)
    overview_p = doc.add_paragraph()
    fmt(overview_p.add_run('Pollution — Reporting Topics, Sub-Topics & ESRS / GRI Alignment'),
        'Pollution — Reporting Topics, Sub-Topics & ESRS / GRI Alignment', sz=10, bold=True)

    t0 = doc.add_table(rows=12, cols=4)
    t0.style = 'Table Grid'
    for j, h in enumerate(['Sub-Topic', 'GRI Reference', 'ESRS (L1–L2)', 'ESRS (L3–L4)']):
        hdr_cell(t0.rows[0].cells[j], h)
    overview_rows = [
        ('Processes to Identify and Assess Material Pollution-Related Impacts, Risks and Opportunities',
         '—', 'ESRS E2 | E2-IRO-1', 'E2-IRO-1-11-(a) | E2-IRO-1-11-(b)'),
        ('Policies Related to Pollution', '—', 'ESRS E2 | E2-1',
         'E2-1-15-(a) | E2-1-15-(b) | E2-1-15-(c) | E2-1-AR-12'),
        ('Actions and Resources Related to Pollution', '—', 'ESRS E2 | E2-2',
         'E2-2-19 | E2-2-AR-13 | E2-2-AR-15'),
        ('Targets Related to Pollution', '—', 'ESRS E2 | E2-3',
         'E2-3-23 | E2-3-24 | E2-3-25 | E2-3-AR-16–18'),
        ('Pollution of Air', 'GRI 305-7', 'ESRS E2 | E2-4',
         'E2-4-28 | E2-4-AR-20 | E2-4-AR-22 | E2-4-30 | E2-4-31'),
        ('Pollution of Water', '—', 'ESRS E2 | E2-4',
         'E2-4-28 | E2-4-AR-20 | E2-4-AR-22 | E2-4-AR-23 | E2-4-30 | E2-4-31'),
        ('Pollution of Soil', '—', 'ESRS E2 | E2-4',
         'E2-4-28 | E2-4-AR-20 | E2-4-AR-22 | E2-4-AR-23 | E2-4-30 | E2-4-31'),
        ('Integrated Pollution Control & BAT Performance Tracking', '—', 'ESRS E2 | E2-4',
         'E2-4-AR-25-(a–e)'),
        ('Substances of Concern', '—', 'ESRS E2 | E2-5', 'E2-5-34'),
        ('Substances of Very High Concern', '—', 'ESRS E2 | E2-5', 'E2-5-35'),
        ('Anticipated Financial Effects from Material Pollution-Related Risks and Opportunities',
         '—', 'ESRS E2 | E2-6',
         'E2-6-39 | E2-6-40 | E2-6-41 | E2-6-AR-31 | E2-6-AR-33'),
    ]
    for i, row in enumerate(overview_rows):
        for j, val in enumerate(row):
            data_cell(t0.rows[i + 1].cells[j], val, row_hdr=(j == 0))

    # ═══════════════════════════════════════════════════════════
    # 5.3.1  Processes to Identify and Assess
    # ═══════════════════════════════════════════════════════════
    heading(doc, '5.3.1  Processes to Identify and Assess Material Pollution-Related '
            'Impacts, Risks and Opportunities', level=2)
    code_para(doc, 'ESRS E2-IRO-1 | ESRS E2-IRO-1-11-(a) | ESRS E2-IRO-1-11-(b)')

    # 5.3.1.1
    heading(doc, '5.3.1.1  Information About the Process to Identify Actual and Potential '
            'Pollution-Related Impacts, Risks and Opportunities', level=3)
    code_para(doc, 'ESRS E2-IRO-1-11-(a)')
    body(doc,
         "XYZ Logistics LLC maintains a formal, structured process to identify and assess material "
         "pollution-related impacts, risks, and opportunities across its operations and value chain. "
         "The Company applies a systematic double materiality assessment methodology, combining an "
         "impact materiality lens — the actual and potential effects of the Company's operations on "
         "ecosystems, water bodies, soils, and communities — with a financial materiality lens "
         "covering risks and opportunities arising from pollution-related regulatory, reputational, "
         "and transition factors. This integrated assessment is conducted annually and forms the "
         "foundation of the Company's environmental reporting and strategic planning cycle.")

    # 5.3.1.2
    heading(doc, '5.3.1.2  Sites and Activities Screened for Pollution-Related Impacts, '
            'Risks and Opportunities — and Methodologies, Assumptions and Tools Used', level=3)
    code_para(doc, 'ESRS E2-IRO-1-11-(a)')
    bool_table(doc,
               'State whether sites and activities were screened for pollution-related '
               'impacts, risks, and opportunities',
               'Yes')
    body(doc,
         "All operating sites and business activities are subject to systematic screening for "
         "pollution-related impacts, risks, and opportunities as part of the annual environmental "
         "risk register review. The screening process applies internationally recognised "
         "methodologies, including ISO 14001-aligned environmental aspect and impact evaluation, "
         "EU Industrial Emissions Directive (IED) applicability assessments, and site-specific "
         "environmental baseline studies where elevated risk is identified. Assumptions "
         "underpinning the screening include the application of conservative emission thresholds, "
         "worst-case scenario modelling for incident probability assessments, and reference to "
         "current EU BAT Conclusions for sector-relevant installations. Geospatial mapping tools "
         "are deployed to identify proximity to sensitive receptors including water bodies, "
         "Natura 2000 sites, and populated areas.")

    # 5.3.1.3
    heading(doc, '5.3.1.3  Stakeholder and Community Consultations Conducted in the '
            'Pollution Impact Assessment Process', level=3)
    code_para(doc, 'ESRS E2-IRO-1-11-(b)')
    bool_table(doc,
               'State whether the undertaking conducted consultations, particularly with '
               'affected communities, in the pollution impact assessment process',
               'Yes')
    body(doc,
         "XYZ Logistics LLC conducted structured consultations with affected communities and "
         "relevant stakeholders as an integral component of its pollution impact assessment "
         "process during FY 2025. Consultations were carried out at all operating sites adjacent "
         "to residential or environmentally sensitive areas. Activities included community liaison "
         "meetings, engagement with local environmental authorities, and formalised grievance "
         "mechanisms accessible to affected parties. The outcomes of these consultations are "
         "documented, reviewed by the Pollution Prevention Council, and incorporated into "
         "site-level environmental management plans. Where material community-level concerns are "
         "identified, corrective actions are tracked through the Company's environmental incident "
         "management system and reported to the Board's Environmental and Sustainability Committee.")

    # ═══════════════════════════════════════════════════════════
    # 5.3.2  Policies Related to Pollution
    # ═══════════════════════════════════════════════════════════
    heading(doc, '5.3.2  Policies Related to Pollution', level=2)
    code_para(doc, 'ESRS E2-1 | ESRS E2-1-15-(a) | ESRS E2-1-15-(b) | ESRS E2-1-15-(c) | ESRS E2-1-AR-12')

    # 5.3.2.1
    heading(doc, '5.3.2.1  Policy Coverage: Mitigation of Negative Impacts on Air, '
            'Water and Soil', level=3)
    code_para(doc, 'ESRS E2-1-15-(a)')
    bool_table(doc,
               'State whether the pollution policy addresses mitigating negative impacts '
               'related to pollution of air, water and soil',
               'Yes')
    body(doc,
         "XYZ Logistics LLC confirms that its Pollution Prevention and Environmental Stewardship "
         "Policy explicitly addresses the mitigation of negative impacts related to pollution of "
         "air, water, and soil across all operating entities and the upstream and downstream value "
         "chain. The Policy establishes clear emission reduction targets, mandatory best available "
         "technique (BAT) adoption requirements for regulated installations, and site-level "
         "environmental management plan obligations. Policy commitments on air, water, and soil "
         "pollution are reviewed annually in conjunction with updated regulatory standards, "
         "monitoring data, and materiality assessment outcomes.")

    # 5.3.2.2
    heading(doc, '5.3.2.2  Policy Coverage: Substitution and Minimisation of Substances '
            'of Concern and Phase-Out of Substances of Very High Concern', level=3)
    code_para(doc, 'ESRS E2-1-15-(b)')
    bool_table(doc,
               'State whether the pollution policy addresses substituting and minimising '
               'use of substances of concern and phasing out substances of very high concern',
               'Yes')
    body(doc,
         "The Company's Policy explicitly addresses the substitution and minimisation of "
         "substances of concern (SoC) and the progressive phase-out of substances of very high "
         "concern (SVHC) in its operations and products. The Company maintains a Restricted "
         "Substances List (RSL) aligned with REACH Regulation Annex XVII and the SVHC Candidate "
         "List, and applies a substitution programme requiring annual review of all SoC and SVHC "
         "in use, with mandatory substitution timelines for SVHC where technically and "
         "economically feasible alternatives exist.")

    # 5.3.2.3
    heading(doc, '5.3.2.3  Policy Coverage: Prevention of and Response to Incidents and '
            'Emergency Situations Resulting in Pollution', level=3)
    code_para(doc, 'ESRS E2-1-15-(c)')
    bool_table(doc,
               'State whether the pollution policy addresses avoiding incidents and emergency '
               'situations resulting in pollution, and controlling or limiting their impact',
               'Yes')
    body(doc,
         "The Policy addresses the prevention of incidents and emergency situations that may "
         "result in pollution, and the containment and limitation of their impact where prevention "
         "is not fully achieved. Emergency response procedures are maintained at all sites, "
         "incorporating spill containment protocols, emergency shutdown procedures, regulatory "
         "notification requirements, and post-incident environmental impact assessments. All "
         "relevant personnel receive mandatory emergency response training. The Company maintains "
         "a centralised incident register, and all pollution-related incidents are subject to root "
         "cause analysis and corrective action tracking.")

    # 5.3.2.4
    heading(doc, '5.3.2.4  Contextual Alignment of Pollution Policy with the '
            'EU Zero Pollution Action Plan', level=3)
    code_para(doc, 'ESRS E2-1-AR-12')
    body(doc,
         "XYZ Logistics LLC's pollution-related policy commitments are explicitly contextualised "
         "against the EU Zero Pollution Action Plan (ZPAP) and its 2030 targets. The Company has "
         "mapped its internal emission reduction targets, SoC substitution commitments, and water "
         "quality protection objectives against the ZPAP's headline goals, including the reduction "
         "of air pollution health impacts, the reduction of microplastic pollution, and the "
         "improvement of water and soil quality. This alignment is reviewed as part of the annual "
         "policy review cycle and disclosed in the Company's Annual Sustainability Report.")

    # ═══════════════════════════════════════════════════════════
    # 5.3.3  Actions and Resources Related to Pollution
    # ═══════════════════════════════════════════════════════════
    heading(doc, '5.3.3  Actions and Resources Related to Pollution', level=2)
    code_para(doc, 'ESRS E2-2 | ESRS E2-2-19 | ESRS E2-2-AR-13 | ESRS E2-2-AR-15')

    # 5.3.3.1
    heading(doc, '5.3.3.1  Layer in Mitigation Hierarchy to Which Actions '
            'Can Be Allocated — Pollution', level=3)
    code_para(doc, 'ESRS E2-2-19')
    body(doc,
         "XYZ Logistics LLC's pollution-related actions are structured in accordance with the "
         "environmental mitigation hierarchy, prioritising: (i) Avoidance — process redesign to "
         "eliminate emission-generating steps, substitution of high-emission materials, and "
         "installation of enclosed handling systems for volatile substances; (ii) Reduction — "
         "deployment of advanced emission abatement technologies, energy efficiency improvements, "
         "and real-time emission monitoring systems; (iii) Restoration — targeted remediation "
         "programmes at historically contaminated sites; and (iv) Compensation — compensatory "
         "environmental measures applied where residual impacts cannot otherwise be addressed. "
         "The majority of FY 2025 pollution actions are allocated to the Avoidance and Reduction "
         "layers of the mitigation hierarchy.")

    # 5.3.3.2
    heading(doc, '5.3.3.2  Pollution Actions Extending to Upstream and Downstream '
            'Value Chain Engagements', level=3)
    code_para(doc, 'ESRS E2-2-AR-13')
    bool_table(doc,
               'State whether pollution-related actions extend to upstream and downstream '
               'value chain engagements',
               'Yes')
    body(doc,
         "Pollution-related actions extend beyond the Company's direct operational boundary to "
         "encompass upstream and downstream value chain engagements. Supplier Environmental "
         "Assessments (SEAs) are conducted for all Tier 1 suppliers assessed as presenting "
         "elevated pollution risk. Downstream, XYZ Logistics LLC engages customers on the "
         "environmental attributes of its products and services, including guidance on responsible "
         "end-of-life management for products containing substances of concern. Value chain "
         "engagement outcomes are tracked through the Company's supplier sustainability programme "
         "and reported annually.")

    # 5.3.3.3
    heading(doc, '5.3.3.3  Layer in Mitigation Hierarchy to Which Resources '
            'Can Be Allocated — Pollution', level=3)
    code_para(doc, 'ESRS E2-2-19')
    body(doc,
         "Financial and human resources allocated to pollution prevention and control are "
         "prioritised in accordance with risk ranking from the annual environmental risk register. "
         "Resource allocations are directed to the Avoidance and Reduction layers of the "
         "mitigation hierarchy as a primary priority, with Restoration-layer resources committed "
         "for active contaminated site remediation programmes. The Company's EHS Capital "
         "Expenditure (CapEx) budget includes dedicated allocations for emission abatement "
         "technology upgrades, site remediation activities, and environmental monitoring "
         "infrastructure. Resource allocation decisions are reviewed semi-annually by the Chief "
         "Sustainability Officer and presented to the Board's Environmental and Sustainability "
         "Committee on an annual basis.")

    # 5.3.3.4
    heading(doc, '5.3.3.4  Site-Level Pollution Prevention Action Plans', level=3)
    code_para(doc, 'ESRS E2-2-AR-15')
    bool_table(doc,
               'State whether pollution prevention action plans have been implemented '
               'at site level',
               'Yes')
    body(doc,
         "Pollution prevention action plans have been implemented at site level across all primary "
         "operating facilities for the FY 2025 reporting period. These plans cover emission "
         "monitoring protocols, substance management procedures, emergency response plans, and "
         "maintenance schedules for abatement equipment. Site-level action plans are reviewed "
         "annually by Regional Environmental Leads and updated to reflect current risk assessments, "
         "regulatory developments, and performance monitoring outcomes. Facilities subject to IED "
         "permitting maintain site-specific environmental improvement programmes aligned with "
         "applicable BAT Conclusions.")

    # ═══════════════════════════════════════════════════════════
    # 5.3.4  Targets Related to Pollution
    # ═══════════════════════════════════════════════════════════
    heading(doc, '5.3.4  Targets Related to Pollution', level=2)
    code_para(doc, 'ESRS E2-3 | ESRS E2-3-23 | ESRS E2-3-24 | ESRS E2-3-25 | '
              'ESRS E2-3-AR-16 | ESRS E2-3-AR-17 | ESRS E2-3-AR-18')

    # 5.3.4.1
    heading(doc, '5.3.4.1  Target: Prevention and Control of Air Pollutants and '
            'Respective Specific Loads', level=3)
    code_para(doc, 'ESRS E2-3-23-(a)')
    bool_table(doc,
               'State whether a target relates to the prevention and control of air '
               'pollutants and respective specific loads',
               'Yes')
    body(doc,
         "XYZ Logistics LLC has established quantitative pollution reduction targets for the "
         "prevention and control of air pollutants and their respective specific loads. "
         "The Company's air quality target for FY 2030 commits to a 40% reduction in total NOx "
         "and SOx emissions relative to the FY 2022 baseline, and a 35% reduction in particulate "
         "matter (PM2.5 and PM10) over the same period. Annual interim milestones are tracked "
         "through the environmental KPI dashboard and reported quarterly to the Board's "
         "Environmental and Sustainability Committee. The FY 2025 interim milestone of a 16% "
         "reduction in total NOx/SOx emissions relative to baseline was achieved, with a reported "
         "17.2% reduction.")

    # 5.3.4.2
    heading(doc, '5.3.4.2  Target: Prevention and Control of Emissions to Water and '
            'Respective Specific Loads', level=3)
    code_para(doc, 'ESRS E2-3-23-(b)')
    bool_table(doc,
               'State whether a target relates to the prevention and control of emissions '
               'to water and respective specific loads',
               'Yes')
    body(doc,
         "The Company has established targets for the prevention and control of emissions to water "
         "and their respective specific loads. The water emission target for FY 2030 commits to a "
         "30% reduction in total suspended solids (TSS) and chemical oxygen demand (COD) loadings "
         "to receiving water bodies relative to the FY 2022 baseline, with a specific sub-target "
         "of zero direct discharge of untreated process water. Interim FY 2025 performance "
         "indicates a 13.5% reduction in TSS and COD loadings. A supplementary target for the "
         "elimination of discharges to water bodies in areas of high water stress has been set "
         "for FY 2027.")

    # 5.3.4.3
    heading(doc, '5.3.4.3  Target: Prevention and Control of Pollution to Soil and '
            'Respective Specific Loads', level=3)
    code_para(doc, 'ESRS E2-3-23-(c)')
    bool_table(doc,
               'State whether a target relates to the prevention and control of pollution '
               'to soil and respective specific loads',
               'Yes')
    body(doc,
         "Targets for the prevention and control of pollution to soil and respective specific "
         "loads are confirmed as operative for FY 2025. These targets are centred on the "
         "elimination of uncontrolled releases of hazardous substances to land, the progressive "
         "remediation of historically contaminated sites, and the reduction of diffuse soil "
         "contamination. The FY 2030 soil pollution target requires all active contaminated site "
         "remediation programmes to have achieved regulatory closure or be at advanced remediation "
         "stage. As of FY 2025, three of seven identified legacy contaminated sites are at "
         "advanced remediation stage, and one has achieved regulatory closure.")

    # 5.3.4.4
    heading(doc, '5.3.4.4  Target: Prevention and Control of Substances of Concern and '
            'Substances of Very High Concern', level=3)
    code_para(doc, 'ESRS E2-3-23-(d)')
    bool_table(doc,
               'State whether a target relates to the prevention and control of substances '
               'of concern and substances of very high concern',
               'Yes')
    body(doc,
         "XYZ Logistics LLC has established targets for the prevention and control of substances "
         "of concern and substances of very high concern across its operations and supply chain. "
         "The SVHC phase-out target for FY 2028 requires the substitution or elimination of all "
         "SVHC for which technically and economically feasible alternatives have been identified "
         "through the annual RSL review. The SoC minimisation target requires a 25% reduction in "
         "total SoC generated, procured, or used in production by FY 2030 relative to the "
         "FY 2022 baseline. As of FY 2025, 4 of 11 identified SVHC substitution projects have "
         "been completed.")

    # 5.3.4.5
    heading(doc, '5.3.4.5  Ecological Thresholds and Entity-Specific Allocations '
            'Considered When Setting Pollution-Related Targets', level=3)
    code_para(doc, 'ESRS E2-3-24 | ESRS E2-3-AR-16-(a)')
    bool_table(doc,
               'State whether ecological thresholds and entity-specific allocations were '
               'taken into consideration when setting pollution-related targets',
               'Yes')
    body(doc,
         "Ecological thresholds — including applicable environmental quality standards (EQS) for "
         "water bodies and air quality limit values under EU Directive 2008/50/EC — have been "
         "incorporated into the target-setting process. Ecological thresholds were identified "
         "through regulatory limit benchmarking, scientific literature review, and consultation "
         "with competent environmental authorities at local, national, and global levels. "
         "Entity-specific emission allocations were determined through a top-down allocation "
         "approach distributing the Company's permitted emission budget across operating sites on "
         "the basis of site throughput, receptor sensitivity, and applicable permit conditions.")

    # 5.3.4.6
    heading(doc, '5.3.4.6  Methodology for Identifying Ecological Thresholds and '
            'Determining Entity-Specific Allocations', level=3)
    code_para(doc, 'ESRS E2-3-24 | ESRS E2-3-AR-16-(a) | ESRS E2-3-AR-16-(b)')
    body(doc,
         "The methodology applied to identify ecological thresholds combines: (i) a regulatory "
         "benchmark approach referencing applicable EU environmental quality standards, IED permit "
         "conditions, and EU BAT Conclusions; (ii) a scientific literature review drawing on "
         "current ecological risk assessment frameworks; and (iii) site-specific environmental "
         "carrying capacity assessments prepared by accredited environmental consultants. "
         "Entity-specific thresholds for each operating site were determined through proportional "
         "allocation modelling, applying site emission intensity ratios adjusted for receptor "
         "sensitivity classifications established through geospatial environmental sensitivity "
         "mapping.")

    # 5.3.4.7
    heading(doc, '5.3.4.7  Responsibility for Respecting Ecological Thresholds, '
            'Mandatory/Voluntary Status, and Substantial Contribution Criteria', level=3)
    code_para(doc, 'ESRS E2-3-24 | ESRS E2-3-AR-16-(c) | ESRS E2-3-25 | ESRS E2-3-AR-17')

    bool_table(doc,
               'State whether the pollution-related target is mandatory (required by '
               'legislation) or voluntary',
               'Mandatory — Required by legislation and voluntary components adopted in '
               'alignment with the EU Zero Pollution Action Plan')

    bool_table(doc,
               'State whether the pollution-related target addresses shortcomings related '
               'to the Substantial Contribution criteria for Pollution Prevention and Control',
               'Yes')

    body(doc,
         "Responsibility for monitoring performance against ecological thresholds rests with "
         "Regional Environmental Leads, with escalation obligations to the Chief Sustainability "
         "Officer where site-level allocations are at risk of being exceeded. Pollution reduction "
         "targets comprise both mandatory targets required by applicable environmental permitting "
         "conditions and EU regulatory frameworks (including IED and REACH), and voluntary "
         "commitments adopted in alignment with the EU Zero Pollution Action Plan. The mandatory "
         "components address specific shortcomings identified under the Substantial Contribution "
         "criteria for Pollution Prevention and Control under the EU Taxonomy Regulation. The "
         "voluntary components represent additional ambition beyond regulatory minimum requirements.")

    # 5.3.4.8
    heading(doc, '5.3.4.8  Pollution Targets Implemented at Site Level', level=3)
    code_para(doc, 'ESRS E2-3-AR-18')
    body(doc,
         "Site-level implementation of all pollution targets is tracked through the EHS incident "
         "and performance management system, with quarterly progress reporting to the Board's "
         "Environmental and Sustainability Committee. Each operating site maintains a site-level "
         "pollution target implementation plan, reviewed annually by Regional Environmental Leads "
         "and updated to reflect current risk assessments, regulatory developments, and monitoring "
         "data outcomes. As of year-end FY 2025, all primary operating sites have operational "
         "site-level pollution target implementation plans in place.")

if __name__ == '__main__':
    d = Document()
    build_part1(d)
    d.save('/home/user/jcode/test_part1.docx')
    print("Part 1 saved.")
