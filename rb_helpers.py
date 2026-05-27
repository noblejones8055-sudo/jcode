#!/usr/bin/env python3
"""
Helper module — Pollution Report 5.3 (Reference-Aligned Edition).

Visual standard reverse-engineered from the supplied reference reports
(Water Management, DEI, Executive Compensation, Compliance):

  • GRI codes  → Blue  0070C0   ESRS codes → Green 00B050   (font size 9)
  • RULE_01  section-header mapping string  : NO brackets, pipe-separated
  • RULE_02  inline / table-title codes      : MANDATORY [ ] at end of text
  • Never italic on any disclosure tag
  • Body narrative          : Arial 11 Black
  • Table cell values       : Arial 10
  • Table header row        : fill D9D9D9, Bold Black
  • Table total / net rows  : fill F3F3F3, Bold Black
  • Boolean / "Whether" data : Indicator | Framework Reference | Status (FY 2025)
"""

from docx.shared import Pt, RGBColor, Inches
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import re

# ── Colours ──────────────────────────────────────────────────
GRI_BLUE    = RGBColor(0x00, 0x70, 0xC0)
ESRS_GREEN  = RGBColor(0x00, 0xB0, 0x50)
BLACK       = RGBColor(0x00, 0x00, 0x00)
HEADER_FILL = 'D9D9D9'
TOTAL_FILL  = 'F3F3F3'

# ── Low-level ────────────────────────────────────────────────
def set_fill(cell, hex_col):
    tcPr = cell._tc.get_or_add_tcPr()
    for s in tcPr.findall(qn('w:shd')):
        tcPr.remove(s)
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_col)
    tcPr.append(shd)

def runfmt(run, text, sz=11, bold=False, italic=False, color=BLACK):
    run.text           = text
    run.font.name      = 'Arial'
    run.font.size      = Pt(sz)
    run.font.bold      = bold
    run.font.italic    = italic
    run.font.color.rgb = color

def _code_colour(token):
    return GRI_BLUE if token.strip().upper().startswith('GRI') else ESRS_GREEN

def _tokens(raw):
    clean = re.sub(r'[\[\]]', '', raw).strip()
    return [t.strip() for t in clean.split('|') if t.strip()]

# ── Disclosure renderers ─────────────────────────────────────
def add_mapping_string(para, raw):
    """RULE_01 — under a major section heading: no brackets, pipe-separated, sz 9."""
    toks = _tokens(raw)
    for i, t in enumerate(toks):
        runfmt(para.add_run(), t, sz=9, color=_code_colour(t))
        if i < len(toks) - 1:
            runfmt(para.add_run(), ' | ', sz=9, color=BLACK)

def add_inline_code(para, raw):
    """RULE_02 — inline at end of narrative / table title: mandatory [ ], sz 9."""
    toks = _tokens(raw)
    runfmt(para.add_run(), ' [', sz=9, color=BLACK)
    for i, t in enumerate(toks):
        runfmt(para.add_run(), t, sz=9, color=_code_colour(t))
        if i < len(toks) - 1:
            runfmt(para.add_run(), ' | ', sz=9, color=BLACK)
    runfmt(para.add_run(), ']', sz=9, color=BLACK)

def add_code_cell(cell, raw):
    """Coloured framework code inside a table cell (no brackets, sz 9)."""
    cell.paragraphs[0].clear()
    if not raw or raw == '—' or raw == '-':
        runfmt(cell.paragraphs[0].add_run(), '—', sz=9, color=BLACK)
        return
    toks = _tokens(raw)
    for i, t in enumerate(toks):
        runfmt(cell.paragraphs[0].add_run(), t, sz=9, color=_code_colour(t))
        if i < len(toks) - 1:
            runfmt(cell.paragraphs[0].add_run(), ' | ', sz=9, color=BLACK)

# ── Headings ─────────────────────────────────────────────────
def heading(doc, text, level):
    """level 1 → 16pt; 2 → 13pt; 3 → 11pt bold (Arial throughout)."""
    sizes = {1: 16, 2: 13, 3: 11}
    h = doc.add_heading('', level=level)
    runfmt(h.add_run(), text, sz=sizes.get(level, 11), bold=True, color=BLACK)
    return h

# ── Paragraphs ───────────────────────────────────────────────
def body(doc, text, code=None, bold=False):
    """Arial 11 Black narrative; optional inline bracketed code (RULE_02)."""
    p = doc.add_paragraph()
    runfmt(p.add_run(), text, sz=11, bold=bold, color=BLACK)
    if code:
        add_inline_code(p, code)
    return p

def section_codes(doc, raw):
    """Standalone RULE_01 mapping string paragraph under a section heading."""
    p = doc.add_paragraph()
    add_mapping_string(p, raw)
    return p

def note(doc, text, code=None):
    """Caption / exclusion note under a table — narrative w/ optional inline code."""
    return body(doc, text, code=code)

# ── Cell writers ─────────────────────────────────────────────
def _hdr_cell(cell, text):
    set_fill(cell, HEADER_FILL)
    cell.paragraphs[0].clear()
    runfmt(cell.paragraphs[0].add_run(), text, sz=10, bold=True, color=BLACK)

def _val_cell(cell, text, bold=False, fill=None):
    if fill:
        set_fill(cell, fill)
    cell.paragraphs[0].clear()
    runfmt(cell.paragraphs[0].add_run(), str(text), sz=10, bold=bold, color=BLACK)

def _is_total(label):
    s = str(label).strip().lower()
    return s.startswith('total') or s.startswith('net ') or 'combined' in s

# ── Tables ───────────────────────────────────────────────────
def data_table(doc, title, code, col_hdrs, rows):
    """Quantitative table: title (bold 10) + inline bracketed code, then grid.
    Header row shaded D9D9D9 bold; Total/Net rows shaded F3F3F3 bold."""
    p = doc.add_paragraph()
    runfmt(p.add_run(), title, sz=10, bold=True, color=BLACK)
    if code:
        add_inline_code(p, code)
    t = doc.add_table(rows=1 + len(rows), cols=len(col_hdrs))
    t.style = 'Table Grid'
    for j, h in enumerate(col_hdrs):
        _hdr_cell(t.rows[0].cells[j], h)
    for i, row in enumerate(rows):
        total = _is_total(row[0])
        for j, v in enumerate(row):
            _val_cell(t.rows[i + 1].cells[j], v,
                      bold=total, fill=(TOTAL_FILL if total else None))
    return t

def status_table(doc, rows):
    """Boolean / 'Whether' data → Indicator | Framework Reference | Status (FY 2025).
    rows = list of (indicator, framework_code_raw, status_text)."""
    t = doc.add_table(rows=1 + len(rows), cols=3)
    t.style = 'Table Grid'
    for j, h in enumerate(['Indicator', 'Framework Reference', 'Status (FY 2025)']):
        _hdr_cell(t.rows[0].cells[j], h)
    for i, (ind, code, status) in enumerate(rows):
        _val_cell(t.rows[i + 1].cells[0], ind)
        add_code_cell(t.rows[i + 1].cells[1], code)
        _val_cell(t.rows[i + 1].cells[2], status)
    return t

def kpi_table(doc, title, code, rows):
    """KPI dashboard → KPI | GRI | ESRS | FY 2025 Data (framework cols coloured)."""
    p = doc.add_paragraph()
    runfmt(p.add_run(), title, sz=10, bold=True, color=BLACK)
    if code:
        add_inline_code(p, code)
    t = doc.add_table(rows=1 + len(rows), cols=4)
    t.style = 'Table Grid'
    for j, h in enumerate(['KPI', 'GRI', 'ESRS', 'FY 2025 Data']):
        _hdr_cell(t.rows[0].cells[j], h)
    for i, (kpi, gri, esrs, val) in enumerate(rows):
        _val_cell(t.rows[i + 1].cells[0], kpi)
        add_code_cell(t.rows[i + 1].cells[1], gri)
        add_code_cell(t.rows[i + 1].cells[2], esrs)
        _val_cell(t.rows[i + 1].cells[3], val)
    return t

def governance_table(doc, rows):
    """Governance Role | Responsibility."""
    t = doc.add_table(rows=1 + len(rows), cols=2)
    t.style = 'Table Grid'
    for j, h in enumerate(['Governance Role', 'Responsibility']):
        _hdr_cell(t.rows[0].cells[j], h)
    for i, (role, resp) in enumerate(rows):
        _val_cell(t.rows[i + 1].cells[0], role, bold=True)
        _val_cell(t.rows[i + 1].cells[1], resp)
    return t

def overview_table(doc, rows):
    """Topic | GRI Reference | ESRS Reference (framework cols coloured)."""
    t = doc.add_table(rows=1 + len(rows), cols=3)
    t.style = 'Table Grid'
    for j, h in enumerate(['Topic', 'GRI Reference', 'ESRS Reference']):
        _hdr_cell(t.rows[0].cells[j], h)
    for i, (topic, gri, esrs) in enumerate(rows):
        _val_cell(t.rows[i + 1].cells[0], topic, bold=True)
        add_code_cell(t.rows[i + 1].cells[1], gri)
        add_code_cell(t.rows[i + 1].cells[2], esrs)
    return t

def content_index_table(doc, rows):
    """Disclosure / Data Point | GRI Standard | ESRS Reference | Section | Status."""
    t = doc.add_table(rows=1 + len(rows), cols=5)
    t.style = 'Table Grid'
    for j, h in enumerate(['Disclosure / Data Point', 'GRI Standard',
                           'ESRS Reference', 'Section', 'Status']):
        _hdr_cell(t.rows[0].cells[j], h)
    for i, (dp, gri, esrs, sec, status) in enumerate(rows):
        _val_cell(t.rows[i + 1].cells[0], dp)
        add_code_cell(t.rows[i + 1].cells[1], gri)
        add_code_cell(t.rows[i + 1].cells[2], esrs)
        _val_cell(t.rows[i + 1].cells[3], sec)
        _val_cell(t.rows[i + 1].cells[4], status)
    return t
