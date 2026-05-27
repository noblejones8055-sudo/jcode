#!/usr/bin/env python3
"""Helper functions for building the corrected Pollution Report Section 5.3."""

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import re

# ── Colour constants ─────────────────────────────────────────
GRI_BLUE   = RGBColor(0, 70, 127)    # GRI codes  → Blue
ESRS_GREEN = RGBColor(0, 128, 0)     # ESRS codes → Green
BLACK      = RGBColor(0, 0, 0)
LG_HEX     = 'D3D3D3'               # Light grey

# ── Low-level helpers ────────────────────────────────────────
def set_bg(cell, hex_col=LG_HEX):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for s in tcPr.findall(qn('w:shd')):
        tcPr.remove(s)
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_col)
    tcPr.append(shd)

def fmt(run, txt, sz=11, bold=False, italic=False, color=None):
    run.text          = txt
    run.font.name     = 'Arial'
    run.font.size     = Pt(sz)
    run.font.bold     = bold
    run.font.italic   = italic
    run.font.color.rgb = color if color else BLACK

def gri_or_esrs(code):
    return GRI_BLUE if code.strip().upper().startswith('GRI') else ESRS_GREEN

# ── Disclosure renderers ─────────────────────────────────────
def disc_r01(para, raw):
    """RULE_01 – sub-section header: no brackets, GRI=Blue, ESRS=Green, size 9."""
    clean = re.sub(r'[\[\]]', '', raw).strip()
    parts = [p.strip() for p in clean.split('|') if p.strip()]
    for i, p in enumerate(parts):
        r = para.add_run(p)
        fmt(r, p, sz=9, color=gri_or_esrs(p))
        if i < len(parts) - 1:
            s = para.add_run(' | ')
            fmt(s, ' | ', sz=9)

def disc_r02(para, raw):
    """RULE_02 – inline/table-title: brackets mandatory, GRI=Blue, ESRS=Green, size 9."""
    clean = re.sub(r'[\[\]]', '', raw).strip()
    parts = [p.strip() for p in clean.split('|') if p.strip()]
    ob = para.add_run('[');  fmt(ob, '[', sz=9)
    for i, p in enumerate(parts):
        r = para.add_run(p); fmt(r, p, sz=9, color=gri_or_esrs(p))
        if i < len(parts) - 1:
            s = para.add_run(' | '); fmt(s, ' | ', sz=9)
    cb = para.add_run(']'); fmt(cb, ']', sz=9)

# ── Paragraph builders ───────────────────────────────────────
def body(doc, text, bold=False):
    """Standard body paragraph – Arial 11 Black."""
    p = doc.add_paragraph()
    r = p.add_run(text)
    fmt(r, text, sz=11, bold=bold)
    return p

def inline_note(doc, text, codes_raw):
    """Paragraph: text + inline RULE_02 code at end."""
    p = doc.add_paragraph()
    r = p.add_run(text + '  ')
    fmt(r, text + '  ', sz=11)
    disc_r02(p, codes_raw)
    return p

def code_para(doc, raw):
    """Standalone RULE_01 disclosure paragraph under a heading."""
    p = doc.add_paragraph()
    disc_r01(p, raw)
    return p

def heading(doc, text, level):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = 'Arial'
    return h

# ── Table builders ───────────────────────────────────────────
def hdr_cell(cell, text):
    set_bg(cell)
    cell.paragraphs[0].clear()
    r = cell.paragraphs[0].add_run(text)
    fmt(r, text, sz=10, bold=True)

def data_cell(cell, text, row_hdr=False):
    if row_hdr:
        set_bg(cell)
        cell.paragraphs[0].clear()
        r = cell.paragraphs[0].add_run(str(text))
        fmt(r, str(text), sz=10, bold=True)
    else:
        cell.paragraphs[0].clear()
        r = cell.paragraphs[0].add_run(str(text))
        fmt(r, str(text), sz=10)

def bool_table(doc, disclosure_text, response):
    """2-column Boolean table per RULE_04."""
    t = doc.add_table(rows=2, cols=2)
    t.style = 'Table Grid'
    hdr_cell(t.rows[0].cells[0], 'Disclosure')
    hdr_cell(t.rows[0].cells[1], 'Response')
    data_cell(t.rows[1].cells[0], disclosure_text, row_hdr=True)
    data_cell(t.rows[1].cells[1], response)
    return t

def tbl(doc, title, codes_raw, col_hdrs, rows, first_col_hdr=True):
    """Full data table: title paragraph + formatted table."""
    # Title paragraph (RULE_02 codes if provided)
    p = doc.add_paragraph()
    r = p.add_run(title + ('  ' if codes_raw else ''))
    fmt(r, title + ('  ' if codes_raw else ''), sz=10, bold=True)
    if codes_raw:
        disc_r02(p, codes_raw)
    # Table
    t = doc.add_table(rows=1 + len(rows), cols=len(col_hdrs))
    t.style = 'Table Grid'
    for j, h in enumerate(col_hdrs):
        hdr_cell(t.rows[0].cells[j], h)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            is_rh = (j == 0 and first_col_hdr)
            data_cell(t.rows[i + 1].cells[j], val, row_hdr=is_rh)
    return t
